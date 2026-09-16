import json
import os
import re
import tempfile
from pathlib import Path

import streamlit as st
from google import genai
from google.genai import types
from docx import Document


MODEL_NAME = "gemini-3.6-flash"
MAX_FILE_SIZE_MB = 20


st.set_page_config(
    page_title="AI Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)


def get_api_key():
    """Read the Gemini API key from Streamlit secrets or an environment variable."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        key = None

    return key or os.getenv("GEMINI_API_KEY")


def extract_docx_text(uploaded_file):
    """Extract readable text from a DOCX resume."""
    document = Document(uploaded_file)
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                parts.append(row_text)

    return "\n".join(parts)


def extract_txt_text(uploaded_file):
    return uploaded_file.getvalue().decode("utf-8", errors="ignore")


def clean_json_text(text):
    """Remove Markdown code fences if Gemini returns JSON inside them."""
    text = text.strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)
    return text.strip()


def analyze_resume(client, uploaded_file, job_description):
    prompt = f"""
You are an expert resume reviewer and ATS optimization specialist.

Analyze the uploaded resume carefully.

IMPORTANT:
- This is an ATS-readiness estimate, NOT a score from a proprietary ATS vendor.
- Do not invent experience, education, certifications, dates, employers, skills, or achievements.
- Base observations only on the resume and the supplied job description.
- If a job description is supplied, evaluate keyword and requirement alignment against it.
- If no job description is supplied, evaluate general ATS readability, structure, keyword quality,
  measurable achievements, section completeness, formatting risk, and professional relevance.
- Be practical and specific.
- Give examples of how to improve wording, but do not fabricate facts.
- A high score must not be awarded merely because the resume looks attractive.

Return ONLY valid JSON matching this structure:

{{
  "ats_score": 0,
  "score_label": "Poor|Needs Improvement|Good|Very Good|Excellent",
  "summary": "short overall summary",
  "score_breakdown": {{
    "keyword_match": 0,
    "format_and_parseability": 0,
    "skills": 0,
    "experience_and_achievements": 0,
    "section_completeness": 0
  }},
  "strengths": [
    "strength 1",
    "strength 2",
    "strength 3"
  ],
  "critical_issues": [
    "issue 1",
    "issue 2"
  ],
  "improvements": [
    {{
      "priority": "High|Medium|Low",
      "area": "area name",
      "problem": "what is wrong or missing",
      "recommendation": "specific action",
      "example": "example wording only when supported by the resume; otherwise say 'Add a factual example from your experience.'"
    }}
  ],
  "keywords_found": ["keyword 1", "keyword 2"],
  "keywords_missing_or_weak": ["keyword 1", "keyword 2"],
  "ats_format_checks": {{
    "standard_section_headings": true,
    "simple_structure": true,
    "contact_information_present": true,
    "dates_consistent": true,
    "bullet_points_used": true,
    "tables_or_complex_layout_risk": false,
    "headers_or_footers_risk": false
  }},
  "recommended_sections": [
    "section 1",
    "section 2"
  ],
  "next_steps": [
    "step 1",
    "step 2",
    "step 3"
  ]
}}

Scoring guidance:
- ats_score must be an integer from 0 to 100.
- keyword_match should reflect job-description alignment when a job description is provided.
- format_and_parseability should reflect likely ATS extraction reliability.
- skills should reflect relevant, clearly stated skills.
- experience_and_achievements should reflect relevance, clarity, measurable impact, and strong action-oriented bullets.
- section_completeness should reflect the presence and quality of standard resume sections.
- Do not treat visual design alone as an ATS advantage.

JOB DESCRIPTION:
{job_description.strip() if job_description.strip() else "No job description supplied. Perform a general ATS-readiness analysis."}
"""

    suffix = Path(uploaded_file.name).suffix.lower()

    if suffix == ".pdf":
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as temp_file:
            temp_file.write(uploaded_file.getvalue())
            temp_path = temp_file.name

        try:
            gemini_file = client.files.upload(
                file=temp_path,
                config={"mime_type": "application/pdf"},
            )
            response = client.models.generate_content(
                model=MODEL_NAME,
                contents=[gemini_file, prompt],
                config=types.GenerateContentConfig(
                    temperature=0.2,
                    response_mime_type="application/json",
                ),
            )
        finally:
            try:
                os.remove(temp_path)
            except OSError:
                pass
    else:
        if suffix == ".docx":
            resume_text = extract_docx_text(uploaded_file)
        else:
            resume_text = extract_txt_text(uploaded_file)

        if not resume_text.strip():
            raise ValueError("The uploaded resume does not contain readable text.")

        response = client.models.generate_content(
            model=MODEL_NAME,
            contents=[prompt, "\nRESUME TEXT:\n", resume_text],
            config=types.GenerateContentConfig(
                temperature=0.2,
                response_mime_type="application/json",
            ),
        )

    raw = clean_json_text(response.text)

    try:
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(
            "Gemini returned an unexpected response format. Please try the analysis again."
        ) from exc


def show_score(score):
    score = max(0, min(100, int(score)))

    if score < 50:
        color = "#dc2626"
    elif score < 70:
        color = "#d97706"
    elif score < 85:
        color = "#2563eb"
    else:
        color = "#16a34a"

    st.markdown(
        f"""
        <div style="text-align:center;padding:18px;border-radius:14px;
                    background:#f8fafc;border:1px solid #e2e8f0;">
            <div style="font-size:14px;color:#64748b;">ATS READINESS SCORE</div>
            <div style="font-size:56px;font-weight:800;color:{color};">{score}/100</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


st.title("📄 AI Resume ATS Analyzer")
st.write(
    "Upload a resume and get an AI-based ATS-readiness score, keyword analysis, "
    "format checks, and actionable improvement recommendations."
)

with st.sidebar:
    st.header("Settings")
    st.info(
        "For the most useful ATS analysis, paste the target job description. "
        "Without it, the app evaluates general ATS readiness."
    )

    if st.button("Clear current analysis"):
        for key in ("analysis", "analyzed_file"):
            st.session_state.pop(key, None)
        st.rerun()

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx", "txt"],
    help="Maximum file size: 20 MB.",
)

job_description = st.text_area(
    "Target Job Description (recommended)",
    height=220,
    placeholder="Paste the complete job description here...",
)

if uploaded_file:
    file_size_mb = len(uploaded_file.getvalue()) / (1024 * 1024)

    if file_size_mb > MAX_FILE_SIZE_MB:
        st.error(f"File is too large. Please upload a file below {MAX_FILE_SIZE_MB} MB.")
        st.stop()

    st.success(f"Ready: {uploaded_file.name} ({file_size_mb:.2f} MB)")

if st.button("🔍 Analyze Resume", type="primary", disabled=uploaded_file is None):
    api_key = get_api_key()

    if not api_key:
        st.error(
            "Gemini API key not found. Add GEMINI_API_KEY to Streamlit Secrets "
            "or set it as an environment variable."
        )
        st.stop()

    try:
        client = genai.Client(api_key=api_key)

        with st.spinner("Gemini is analyzing the resume..."):
            result = analyze_resume(client, uploaded_file, job_description)

        st.session_state.analysis = result
        st.session_state.analyzed_file = uploaded_file.name

    except Exception as exc:
        st.error(f"Analysis failed: {exc}")

if "analysis" in st.session_state:
    result = st.session_state.analysis

    st.divider()
    st.subheader(f"Analysis: {st.session_state.get('analyzed_file', 'Resume')}")

    col1, col2 = st.columns([1, 2])

    with col1:
        show_score(result.get("ats_score", 0))

    with col2:
        st.markdown(f"### {result.get('score_label', 'ATS Result')}")
        st.write(result.get("summary", ""))

    st.divider()

    st.subheader("📊 Score Breakdown")
    breakdown = result.get("score_breakdown", {})

    breakdown_cols = st.columns(5)
    breakdown_items = [
        ("Keyword Match", "keyword_match"),
        ("Format", "format_and_parseability"),
        ("Skills", "skills"),
        ("Experience", "experience_and_achievements"),
        ("Sections", "section_completeness"),
    ]

    for col, (label, key) in zip(breakdown_cols, breakdown_items):
        with col:
            st.metric(label, f"{breakdown.get(key, 0)}/100")

    st.divider()

    left, right = st.columns(2)

    with left:
        st.subheader("✅ Strengths")
        for item in result.get("strengths", []):
            st.markdown(f"- {item}")

        st.subheader("🔑 Keywords Found")
        keywords = result.get("keywords_found", [])
        st.write(", ".join(keywords) if keywords else "No strong keywords identified.")

    with right:
        st.subheader("⚠️ Critical Issues")
        for item in result.get("critical_issues", []):
            st.markdown(f"- {item}")

        st.subheader("🔎 Missing / Weak Keywords")
        missing = result.get("keywords_missing_or_weak", [])
        st.write(", ".join(missing) if missing else "No major missing keywords identified.")

    st.divider()

    st.subheader("🛠️ Recommended Improvements")

    for index, item in enumerate(result.get("improvements", []), start=1):
        priority = item.get("priority", "Medium")
        with st.expander(
            f"{index}. {item.get('area', 'Improvement')} — {priority} Priority",
            expanded=priority == "High",
        ):
            st.markdown(f"**Problem:** {item.get('problem', '')}")
            st.markdown(f"**Recommendation:** {item.get('recommendation', '')}")
            st.markdown(f"**Example:** {item.get('example', '')}")

    st.divider()

    st.subheader("📋 ATS Format Checks")
    checks = result.get("ats_format_checks", {})

    check_labels = {
        "standard_section_headings": "Standard section headings",
        "simple_structure": "Simple structure",
        "contact_information_present": "Contact information present",
        "dates_consistent": "Dates are consistent",
        "bullet_points_used": "Bullet points used",
        "tables_or_complex_layout_risk": "Tables / complex-layout risk",
        "headers_or_footers_risk": "Headers / footers risk",
    }

    for key, label in check_labels.items():
        value = checks.get(key)
        if value is True:
            st.success(f"✓ {label}")
        elif value is False:
            st.error(f"✗ {label}")
        else:
            st.warning(f"? {label}: Not determined")

    st.divider()

    col1, col2 = st.columns(2)

    with col1:
        st.subheader("📌 Recommended Sections")
        for item in result.get("recommended_sections", []):
            st.markdown(f"- {item}")

    with col2:
        st.subheader("🚀 Next Steps")
        for item in result.get("next_steps", []):
            st.markdown(f"- {item}")

st.caption(
    "ATS score is an AI-generated resume-readiness estimate. It is not an official score "
    "from a specific Applicant Tracking System."
)
